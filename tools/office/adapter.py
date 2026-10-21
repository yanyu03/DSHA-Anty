#!/usr/bin/env python3
"""Office Tool: inspect local macro-free OOXML, avoid Internet and macros."""
from __future__ import annotations
import argparse, importlib.util, json, pathlib, stat, zipfile

EXTS={'.docx':'docx','.xlsx':'xlsx','.pptx':'pptx'}
MAX_BYTES=12_000_000;MAX_EXPANDED=60_000_000;MAX_ENTRIES=1500

def inspect(value):
    path=pathlib.Path(value)
    if not path.exists() or not stat.S_ISREG(path.lstat().st_mode):raise ValueError('ERR_INVALID_OR_LINKED_FILE')
    typ=EXTS.get(path.suffix.lower())
    if typ is None:raise ValueError('ERR_UNSUPPORTED_OR_MACRO_FORMAT')
    if path.stat().st_size > MAX_BYTES:raise ValueError('ERR_TOO_LARGE')
    try:
        with zipfile.ZipFile(path) as z:
            infos=z.infolist()
            if len(infos)>MAX_ENTRIES or sum(x.file_size for x in infos)>MAX_EXPANDED:
                raise ValueError('ERR_ARCHIVE_BOMB_LIMIT')
            if any(x.filename.startswith('..') or x.filename.startswith('/') for x in infos):
                raise ValueError('ERR_INVALID_ZIP_ENTRY')
            names={x.filename for x in infos}
            if any('vbaproject.bin' in x.lower() for x in names):raise ValueError('ERR_EMBEDDED_MACRO')
            if '[Content_Types].xml' not in names:raise ValueError('ERR_NOT_OOXML')
    except zipfile.BadZipFile:raise ValueError('ERR_BAD_ARCHIVE')
    if typ=='docx':
        from docx import Document
        document=Document(str(path))
        result={'paragraphs':len(document.paragraphs),'tables':len(document.tables),
                'headings':[p.text[:100] for p in document.paragraphs if p.style.name.startswith('Heading')][:60]}
    elif typ=='xlsx':
        from openpyxl import load_workbook
        book=load_workbook(str(path),read_only=True,data_only=False,keep_links=False)
        try:result={'sheets':book.sheetnames[:60],'sheet_count':len(book.sheetnames)}
        finally:book.close()
    else:
        from pptx import Presentation
        ppt=Presentation(str(path));result={'slides':len(ppt.slides)}
    return {'status':'PASS','tool':'office','kind':typ,'summary':result,
            'macros_executed':False,'network_used':False}

def main(argv=None):
    p=argparse.ArgumentParser();p.add_argument('action',choices=['status','inspect']);p.add_argument('file',nargs='?');a=p.parse_args(argv)
    if a.action=='status':
        present={name:bool(importlib.util.find_spec(name)) for name in ('docx','openpyxl','pptx')}
        out={'status':'AVAILABLE' if all(present.values()) else 'DEPENDENCY_MISSING',
             'tool':'office','dependencies':present,'side_effects':'none'}
    else:
        try:
            if not a.file:raise ValueError('ERR_FILE_REQUIRED')
            out=inspect(a.file)
        except (OSError,ValueError,KeyError,ImportError,RuntimeError,IndexError) as e:
            out={'status':'ERR_BLOCK','reason':str(e)[:120],'network_used':False}
    print(json.dumps(out,ensure_ascii=False))
    return 0 if out['status'] in ('PASS','AVAILABLE') else 2
if __name__=='__main__':raise SystemExit(main())
